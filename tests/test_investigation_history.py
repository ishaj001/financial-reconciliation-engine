import pandas as pd
from openpyxl import load_workbook

from src.reconciliation import reconcile
from src.database import save_run, table, record_manual_match, run_comparison, start_run, fail_run
from src.investigation import candidate_matches, source_totals_by_currency, matched_controls_by_currency, unresolved_controls_by_currency
from src.reporting import excel_bytes

COLS=["transaction_id","transaction_date","settlement_date","amount","currency","transaction_type","status","source_system"]
def ledger(rows): return pd.DataFrame(rows,columns=COLS)
def item(key,side,amount=10,date="2025-02-01"):
    return [key,date,"2025-02-02",amount,"USD","purchase","settled",side]

def test_run_history_comparison_currency_totals_and_excel(tmp_path):
    path=tmp_path/"runs.sqlite"
    a1=ledger([item("a1","A"),item("a2","A",15)])
    b1=ledger([item("a1","B"),item("b2","B",15)])
    out1=reconcile(a1,b1); run1=save_run(out1,a1,b1,path=path)
    a2=ledger([item("a1","A"),item("a2","A",15)])
    b2=ledger([item("a1","B"),item("a2","B",15)])
    out2=reconcile(a2,b2); run2=save_run(out2,a2,b2,path=path)
    compared=run_comparison(run1,run2,path)
    assert "outcome" in compared.section.tolist()
    assert "gross_source_difference" in compared.section.tolist()
    totals=source_totals_by_currency(run2,path)
    assert totals.system_a_records.sum()==2 and totals.system_b_records.sum()==2 and set(totals.currency)=={"USD"}
    matched=matched_controls_by_currency(run2,path)
    assert matched.iloc[0].matched_keys==2
    assert unresolved_controls_by_currency(run1,path).empty
    assert table(path,"SELECT COUNT(*) n FROM reconciliation_results WHERE run_id=?",(run1,)).iloc[0].n==3
    workbook=load_workbook(__import__("io").BytesIO(excel_bytes(run2,path)),read_only=True)
    assert {"Executive Summary","Control Totals","Matched Controls","Unresolved Controls","Matched","Unmatched","Exceptions","Resolution History","Manual Decisions"}.issubset(set(workbook.sheetnames))
    assert workbook["Control Totals"].max_row>=2
    assert workbook["Control Totals"]["G2"].value==0

def test_failed_run_is_audited(tmp_path):
    path=tmp_path/"failed.sqlite"
    run=start_run(("bad-a.csv","bad-b.csv"),{},path)
    fail_run(run,"Malformed input",path)
    saved=table(path,"SELECT processing_status FROM reconciliation_runs WHERE run_id=?",(run,))
    assert saved.iloc[0].processing_status=="failed"
    assert table(path,"SELECT event_type FROM run_events WHERE run_id=? ORDER BY event_id",(run,)).event_type.tolist()==["started","failed"]

def test_gross_and_outcome_totals_stay_separate_by_currency(tmp_path):
    from src.analytics import outcome_summary
    path=tmp_path/"currency.sqlite"
    a=ledger([item("usd","A",10),["eur","2025-02-01","2025-02-02",5,"EUR","purchase","settled","A"]])
    b=ledger([item("usd","B",10),["eur","2025-02-01","2025-02-02",4,"EUR","purchase","settled","B"]])
    output=reconcile(a,b); run=save_run(output,a,b,path=path)
    totals=source_totals_by_currency(run,path).set_index("currency")
    assert totals.loc["USD","gross_amount_difference"]==0
    assert totals.loc["EUR","gross_amount_difference"]==1
    outcomes=outcome_summary(run,path)
    assert set(outcomes.currency)=={"USD","EUR"}
    assert outcomes.loc[outcomes.outcome.eq("mismatch"),"currency"].tolist()==["EUR"]

def test_candidate_suggestions_are_ambiguous_and_need_an_audited_decision(tmp_path):
    path=tmp_path/"candidates.sqlite"
    a=ledger([item("A-MISSING","A",10)])
    b=ledger([item("B-CANDIDATE-1","B",10.1),item("B-CANDIDATE-2","B",10.2)])
    output=reconcile(a,b); run=save_run(output,a,b,path=path)
    suggestions=candidate_matches(run,"A-MISSING",path)
    assert len(suggestions)==2 and all(x["ambiguous"] for x in suggestions)
    assert table(path,"SELECT COUNT(*) n FROM manual_match_decisions").iloc[0].n==0
    first=suggestions[0]
    record_manual_match(run,"A-MISSING",first["system_a_row"],first["system_b_row"],"confirmed","Analyst verified settlement reference",path)
    decision=table(path,"SELECT decision,justification FROM manual_match_decisions WHERE run_id=?",(run,))
    assert decision.iloc[0].decision=="confirmed"
    assert table(path,"SELECT outcome FROM reconciliation_results WHERE run_id=? AND transaction_id='A-MISSING'",(run,)).iloc[0].outcome=="missing_in_b"
